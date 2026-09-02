"""Run classify_signal() on experiment.xlsx inputs and write actual categories."""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from scrapers.classifier import classify_signal  # noqa: E402

XLSX = ROOT / "experiment.xlsx"
INPUT_COL = 4  # "3. System Input"
ACTUAL_COL = 5  # "4. Actual Categories"
MATCH_COL = 6
EXPECTED_COL = 3


def _format_categories(categories: list[str]) -> str:
    return ", ".join(categories) if categories else "(none)"


def _normalize(label: str) -> set[str]:
    text = (label or "").strip()
    if not text or text == "(none)":
        return set()
    return {part.strip() for part in text.split(",") if part.strip()}


def _score(expected: str, actual: str) -> str:
    exp = _normalize(expected)
    act = _normalize(actual)
    if exp == act:
        return "exact"
    if exp and exp <= act:
        return "partial+"
    if exp & act:
        return "partial"
    return "miss"


def run(*, write: bool = True) -> list[dict]:
    wb = load_workbook(XLSX)
    ws = wb["Experiment"]
    results: list[dict] = []

    for row in range(2, ws.max_row + 1):
        text = ws.cell(row=row, column=INPUT_COL).value
        expected = ws.cell(row=row, column=EXPECTED_COL).value or ""
        if not text:
            continue
        actual = _format_categories(classify_signal(str(text)).categories)
        match = _score(str(expected), actual)
        results.append(
            {
                "row": row,
                "expected": expected,
                "actual": actual,
                "match": match,
            }
        )
        if write:
            ws.cell(row=row, column=ACTUAL_COL, value=actual)
            ws.cell(row=row, column=MATCH_COL, value=match)

    if write:
        wb.save(XLSX)
        print(f"Updated {XLSX}")
    return results


if __name__ == "__main__":
    for item in run(write="--dry-run" not in sys.argv):
        print(
            f"row {item['row']}: expected={item['expected']!r} "
            f"actual={item['actual']!r} ({item['match']})"
        )
