"""Generate experiment.xlsx for category-classification accuracy testing."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "experiment.xlsx"

CATEGORIES = [
    "potholes",
    "noise",
    "sanitation",
    "violent_crime",
    "property_crime",
    "traffic_safety",
    "emergencies",
    "public_safety",
    "housing",
    "immigration",
]

# 20 held-out examples: mix of categories, negatives, multi-label, and edge cases.
CASES: list[tuple[str, str]] = [
    (
        "Hit a huge dip on Irvine Blvd and my rim is bent",
        "potholes",
    ),
    (
        "Neighbors blasting music until 2am again, third weekend in a row",
        "noise",
    ),
    (
        "Someone dumped a couch and a mattress behind the plaza again",
        "sanitation",
    ),
    (
        "Someone pulled a gun during an argument at the park, everybody ran",
        "violent_crime",
    ),
    (
        "Catalytic converter stolen off my car overnight, second one on our street",
        "property_crime",
    ),
    (
        "Almost got hit in the crosswalk again, drivers do not even slow down",
        "traffic_safety",
    ),
    (
        "They are telling the whole neighborhood to pack up and be ready to leave",
        "emergencies",
    ),
    (
        "Street lamp by the trail has been out for a month, pitch black at night",
        "public_safety",
    ),
    (
        "Our landlord just raised the lease renewal by 400 a month",
        "housing",
    ),
    (
        "ICE agents were seen outside the market on 1st Street this morning",
        "immigration",
    ),
    (
        "Best boba spot in Irvine? Just moved here",
        "(none)",
    ),
    (
        "Three bed two bath condo available for lease near the park, pool and gym",
        "(none)",
    ),
    (
        "Gas leak forced us to evacuate, fire trucks and police blocking the whole street",
        "emergencies, public_safety",
    ),
    (
        "Street racers doing donuts at midnight, whole neighborhood called the cops",
        "noise, public_safety",
    ),
    (
        "Landlord will not fix the mold but raised the rent anyway",
        "housing",
    ),
    (
        "March down Alton today, streets closed around city hall",
        "public_safety",
    ),
    (
        "Graffiti on the sound wall keeps coming back the day after the city paints over it",
        "sanitation",
    ),
    (
        "My neighbor was detained even though his papers were in process",
        "immigration",
    ),
    (
        "pothole on Culver wrecked my rim",
        "potholes",
    ),
    (
        "The scariest thing around here is the cost of living honestly",
        "(none)",
    ),
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")


def _style_header_row(ws, row: int, cols: int) -> None:
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _autosize(ws, widths: dict[int, float]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def build_workbook() -> Workbook:
    wb = Workbook()

    # --- Instructions ---
    ws_info = wb.active
    ws_info.title = "Instructions"
    info_lines = [
        ("CivicPulse Category Accuracy Experiment", True),
        ("", False),
        ("Goal", True),
        (
            "Measure how accurately the signal classifier assigns civic issue categories "
            "to short resident-style posts (signals).",
            False,
        ),
        ("", False),
        ("Workflow (4 columns on the Experiment sheet)", True),
        ("1. Input", "Write or review the 20 signal texts to classify."),
        ("2. Expected Categories", "Human ground truth — what categories should apply."),
        ("3. System Input", "Copy of the text fed into classify_signal() (pre-filled)."),
        ("4. Actual Categories", "Classifier output — fill in when you run the experiment."),
        ("", False),
        ("How to run", True),
        (
            "From the repo root: python -c \"from scrapers.classifier import classify_signal; "
            "r=classify_signal('YOUR TEXT'); print(', '.join(r.categories) or '(none)')\"",
            False,
        ),
        (
            "Or use: python scripts/run_category_experiment.py (after filling column 4 manually "
            "or to batch-run and compare).",
            False,
        ),
        ("", False),
        ("Scoring", True),
        ("Exact match", "Actual categories exactly match expected (order ignored)."),
        ("Partial credit", "At least one expected category present with no false positives."),
        ("Miss", "No overlap with expected, or extra wrong categories."),
        ("", False),
        ("Notes", True),
        ("(none)", "Use when no civic category should apply."),
        ("Multi-label", "Separate multiple expected categories with commas."),
    ]
    for row_idx, item in enumerate(info_lines, start=1):
        if isinstance(item[1], bool):
            ws_info.cell(row=row_idx, column=1, value=item[0]).font = Font(bold=True, size=12)
        else:
            ws_info.cell(row=row_idx, column=1, value=item[0]).font = Font(bold=True)
            ws_info.cell(row=row_idx, column=2, value=item[1]).alignment = WRAP
    _autosize(ws_info, {1: 28, 2: 90})

    # --- Categories reference ---
    ws_cat = wb.create_sheet("Categories")
    ws_cat.append(["Category", "Description"])
    _style_header_row(ws_cat, 1, 2)
    descriptions = {
        "potholes": "Road damage, pavement cracks, street repair",
        "noise": "Loud neighbors, construction noise, fireworks, barking",
        "sanitation": "Trash, illegal dumping, sewage, rats, graffiti",
        "violent_crime": "Shootings, assault, robbery, stabbings",
        "property_crime": "Theft, break-ins, vandalism, porch pirates",
        "traffic_safety": "Crashes, speeding, crosswalks, reckless driving",
        "emergencies": "Fires, floods, evacuations, power outages, hazmat",
        "public_safety": "General crime/police, protests, missing persons, unsafe areas",
        "housing": "Rent burden, evictions, homelessness, landlord issues",
        "immigration": "ICE activity, deportation, asylum, border enforcement",
    }
    for cat in CATEGORIES:
        ws_cat.append([cat, descriptions[cat]])
    for row in ws_cat.iter_rows(min_row=2, max_row=ws_cat.max_row, min_col=1, max_col=2):
        for cell in row:
            cell.alignment = WRAP
    _autosize(ws_cat, {1: 18, 2: 70})

    # --- Main experiment sheet ---
    ws = wb.create_sheet("Experiment")
    headers = [
        "#",
        "1. Input",
        "2. Expected Categories",
        "3. System Input",
        "4. Actual Categories",
        "Match?",
        "Notes",
    ]
    ws.append(headers)
    _style_header_row(ws, 1, len(headers))

    for idx, (text, expected) in enumerate(CASES, start=1):
        ws.append([idx, text, expected, text, "", "", ""])
        row = ws.max_row
        for col in (2, 3, 4, 7):
            ws.cell(row=row, column=col).alignment = WRAP

    ws.freeze_panes = "A2"
    _autosize(ws, {1: 5, 2: 55, 3: 22, 4: 55, 5: 22, 6: 10, 7: 25})

    return wb


def main() -> None:
    build_workbook().save(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
